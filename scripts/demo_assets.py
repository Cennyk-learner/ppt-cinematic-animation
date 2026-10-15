# 生成一套不依赖生图工具的演示素材，配合 build_page_example.ps1 把"合成 → 组装 → 导出核对"整条链路跑通。
# 换新机器、或者改了工具函数以后，先跑这一对脚本确认环境没问题，再做正式的页面。
# 用法: python demo_assets.py <工作目录>        素材写到 <工作目录>/out/
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from render_tools import *
import pptx
from pptx.util import Emu

work = sys.argv[1] if len(sys.argv) > 1 else "."
out = os.path.join(work, "out"); os.makedirs(out, exist_ok=True)
J = lambda name: os.path.join(out, name)

# 底稿：一页空白的 16:9
p = pptx.Presentation(); p.slide_width = Emu(12192000); p.slide_height = Emu(6858000)
p.slides.add_slide(p.slide_layouts[6]); p.save(J("base.pptx"))


def scene(seed, tint):
    """用噪声和渐变冒充一张夜景底图"""
    r = np.random.default_rng(seed)
    n = cv2.resize(r.random((6, 10)).astype(np.float32), (W, H), interpolation=cv2.INTER_CUBIC)
    ridge = (Y > H * (0.45 + 0.12 * np.sin(X / 260.0 + seed) + 0.05 * n)).astype(np.float32)
    img = (0.10 + 0.10 * n)[..., None] * np.float32(tint) * (1 - ridge[..., None]) + (0.05 + 0.06 * n)[..., None] * np.float32([0.6, 0.8, 0.7]) * ridge[..., None]
    return blur(img, 2)


a, b = scene(1, [1.0, 0.85, 0.7]), scene(4, [0.8, 0.9, 1.0])
rain = Rain(500, seed=3, angle=-0.18)
lamp = pool(W * 0.62, H * 0.52, 60)[..., None] * np.float32([0.2, 0.9, 0.3])
# 第一幕：下雨的循环动图，坡上一盏灯在闪
loop_gif(J("scene_a.gif"), lambda i, t: a + rain.draw(t)[..., None] * 0.3 + lamp * blink(i, 20, 1, 0.3), n=20, nx=5, colors=128)
# 第二幕：静态底图 + 一条交给原生擦除"画出来"的线
still(J("scene_b.png"), b)
xs = np.linspace(0, 1, 200)
pts = np.stack([160 + xs * 1100, 520 - 260 * (0.2 * xs + 0.8 * xs ** 2.6)], 1)
putpos(J("pos.json"), "line", wrgba(J("line.png"), np.float32([60, 190, 255]) / 255, glow(line_layer(pts, 5), 9)))
putpos(J("pos.json"), "mark", wrgba(J("mark.png"), np.float32([45, 45, 235]) / 255, glow(line_layer([(820, 160), (820, 560)], 3, 12, 9), 5, 0.6)))
make_band(J("band.png")); make_fog(J("fog.png"))
print("demo assets ok ->", out)
