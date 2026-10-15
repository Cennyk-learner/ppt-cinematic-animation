# 灯塔演示页的素材合成。三张底图（阴天底片 / 同机位雨夜 / 同机位清晨）由 prompt_gen.txt 一次生成，已放在 gen/ 里。
# 用法:  python render.py <工作目录>          素材写到 <工作目录>/out/，再用 build.ps1 组装成 PPT
# 依赖:  numpy opencv-python-headless pillow python-pptx，ffmpeg 在 PATH 里，Ma Shan Zheng 字体（Google Fonts，免费）
import sys, os, glob, math
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "scripts"))
from render_tools import *
import pptx
from pptx.util import Emu

work = sys.argv[1] if len(sys.argv) > 1 else HERE
out = os.path.join(work, "out"); os.makedirs(out, exist_ok=True)
J = lambda name: os.path.join(out, name)
G = lambda name: load(os.path.join(HERE, "gen", name))


def find_font(names):
    roots = [os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\Windows\Fonts"), r"C:\Windows\Fonts"]
    for r in roots:
        for n in names:
            hit = glob.glob(os.path.join(r, n))
            if hit:
                return hit[0]
    raise SystemExit("找不到书法字体：请安装 Ma Shan Zheng（fonts.google.com/specimen/Ma+Shan+Zheng）")


BRUSH = find_font(["MaShanZheng*.ttf"])

# 底稿：一页空白 16:9
p = pptx.Presentation(); p.slide_width = Emu(12192000); p.slide_height = Emu(6858000)
p.slides.add_slide(p.slide_layouts[6]); p.save(J("base.pptx"))

night, dawn = G("night.jpg"), G("dawn.jpg")
LX, LY = 1405.0, 321.0                     # 灯笼在 1920x1080 画面里的位置（从夜景图最亮处找出来的）
WARM = np.float32([0.62, 0.86, 1.0])       # BGR，暖白灯光
rain = Rain(650, seed=3, angle=-0.20)
rain2 = Rain(320, seed=9, angle=-0.26, speed=(1000, 1500), length=(12, 30))
WORD = text_mask("一束光", BRUSH, 190, 170, 175, angle=3)   # 烧进动图、被光扫到才亮的字
WX, WY = 470.0, 280.0                                      # 这组字的大致中心
N, DT = 40, 0.08                                           # 40 帧一圈 = 光束转一周 3.2 秒


# ---------- 第一幕：雨夜，灯塔的光束一圈圈扫过，扫过天上那组字时字被照亮 ----------
def night_frame(i, t):
    ph = 2 * math.pi * i / N
    c, s = math.cos(ph), math.sin(ph)            # c<0 光束朝左扫海面，c>0 朝右（照向崖后），s<0 朝向镜头
    reach = 1650 * abs(c) + 120
    dst = (LX + reach * (1 if c > 0 else -1), LY - 70 * abs(c) + 25)
    side = min(1.0, abs(c) / 0.55) ** 1.5         # 光束转到正对/背对镜头时，侧面看到的光柱消失
    beam = beam_cone((LX, LY), dst, w_src=8, w_dst=70 + 120 * abs(s), soft=30) * (0.45 + 0.55 * haze(t, 70)) * side
    toward = max(0.0, -s)                        # 朝向镜头时灯笼刺眼
    fr = night * 0.92
    fr = fr + beam[..., None] * WARM * 0.34
    rl = rain.draw(t) * 0.55 + blur(rain2.draw(t), 1.1) * 0.35
    fr = fr + (rl * (0.12 + 1.6 * beam))[..., None] * np.float32([0.9, 0.95, 1.0])
    fl = pool(LX, LY, 14 + 16 * toward) * (0.55 + 0.9 * toward) + pool(LX, LY, 140) * 0.14 * (0.4 + toward)
    fr = fr + fl[..., None] * WARM
    # 光束扫过时把字照亮；用光束在字中心的强度决定亮度
    g = float(np.clip(beam[int(WY), int(WX)] * 2.6, 0, 1))
    fr = lit_text(fr, WORD, 0.16 + 0.84 * g, color=(0.80, 0.92, 1.0))
    return fr


loop_gif(J("night.gif"), night_frame, n=N, nx=6, dt=DT, colors=200)
still(J("_night_still.png"), night_frame(N * 3 // 8, 10))

# ---------- 第三幕：清晨，太阳倒影里的碎光在闪 ----------
sea = ((Y > 575) & (Y < 840)).astype(np.float32) * np.clip(1 - np.abs(X - 330) / 520, 0, 1)
gl_lum = np.clip((dawn.mean(2) - 0.55) * 3, 0, 1) * sea


def dawn_frame(i, t):
    a = haze(t * 3.0, 260); b = haze(t * 2.3 + 5, -310)
    sp = np.clip((a * b - 0.30) * 9, 0, 1) * gl_lum
    return dawn + blur(sp, 1.2)[..., None] * np.float32([0.75, 0.9, 1.0]) * 0.8


loop_gif(J("dawn.gif"), dawn_frame, n=24, nx=6, dt=0.08, colors=220)

# ---------- 第三幕的航线：渔船 → 灯塔脚下，金色虚线，交给原生擦除"画"出来 ----------
route = chaikin([(470, 612), (700, 640), (930, 655), (1120, 650), (1250, 640)], 4)
putpos(J("pos.json"), "route", wrgba(J("route.png"), np.float32([70, 200, 255]) / 255, glow(line_layer(route, 4, dash=16, gap=11), 6, 0.7)))
putpos(J("pos.json"), "route_pts", [[round(float(x) / 2, 1), round(float(y) / 2, 1)] for x, y in route[::6]])
dot = pool(30, 30, 7)[:60, :60]
dot_a = np.clip(dot * 1.6, 0, 1)
cv2.imencode(".png", (np.dstack([np.ones((60, 60, 3)) * np.float32([0.75, 0.93, 1.0]), dot_a]) * 255).astype(np.uint8))[1].tofile(J("dot.png"))

make_band(J("band.png")); make_fog(J("fog.png"))
print("lighthouse assets ->", out)
